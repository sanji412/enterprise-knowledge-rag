#!/usr/bin/env python3
"""Build the versioned enterprise evaluation corpus from Markdown sources."""

from __future__ import annotations

import hashlib
import json
import re
import shutil
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path

import pypdf as PyPDF2
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib.colors import HexColor
from reportlab.lib.pagesizes import LETTER
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

ROOT = Path(__file__).resolve().parents[1]
SOURCE_DIR = ROOT / "evals" / "corpus" / "source"
GENERATED_DIR = ROOT / "evals" / "corpus" / "generated"
DOCX_BODY_FONT = "Arial Unicode MS"
PDF_CID_FONT = "STSong-Light"
PDF_VISIBLE_FONT = "EnterpriseCJK"
PDF_VISIBLE_FONT_CANDIDATES = (
    Path("/System/Library/Fonts/Supplemental/Arial Unicode.ttf"),
    Path("/usr/share/fonts/truetype/noto/NotoSansCJK-Regular.ttf"),
    Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"),
)
H2_RE = re.compile(
    r"^##\s+(.+?)\s+\{#([A-Za-z0-9][A-Za-z0-9_.:-]*)\}\s*$",
)


@dataclass(frozen=True)
class MarkdownDocument:
    title: str
    introduction: str
    sections: tuple["MarkdownSection", ...]


@dataclass(frozen=True)
class MarkdownSection:
    title: str
    anchor: str
    body: str


def _sha256_bytes(content: bytes) -> str:
    return hashlib.sha256(content).hexdigest()


def _normalize_text(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def _normalized_sha256(text: str) -> str:
    return _sha256_bytes(_normalize_text(text).encode("utf-8"))


def _parse_markdown(path: Path) -> MarkdownDocument:
    title = ""
    intro_lines: list[str] = []
    sections: list[MarkdownSection] = []
    section_title: str | None = None
    section_anchor: str | None = None
    body_lines: list[str] = []

    def flush_section() -> None:
        nonlocal body_lines
        if section_title is None or section_anchor is None:
            return
        sections.append(
            MarkdownSection(
                title=section_title,
                anchor=section_anchor,
                body="\n".join(body_lines).strip(),
            ),
        )
        body_lines = []

    for line in path.read_text(encoding="utf-8").splitlines():
        if line.startswith("# ") and not title:
            title = line[2:].strip()
            continue
        match = H2_RE.match(line)
        if match:
            flush_section()
            section_title = match.group(1).strip()
            section_anchor = match.group(2)
            continue
        if section_title is None:
            intro_lines.append(line)
        else:
            body_lines.append(line)
    flush_section()

    if not title or not sections:
        raise ValueError(f"Invalid enterprise source document: {path}")
    return MarkdownDocument(
        title=title,
        introduction="\n".join(intro_lines).strip(),
        sections=tuple(sections),
    )


def _register_pdf_fonts() -> tuple[str, str]:
    if PDF_CID_FONT not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(UnicodeCIDFont(PDF_CID_FONT))
    if PDF_VISIBLE_FONT not in pdfmetrics.getRegisteredFontNames():
        font_path = next(
            (candidate for candidate in PDF_VISIBLE_FONT_CANDIDATES if candidate.is_file()),
            None,
        )
        if font_path is None:
            return PDF_CID_FONT, PDF_CID_FONT
        pdfmetrics.registerFont(TTFont(PDF_VISIBLE_FONT, str(font_path)))
    return PDF_VISIBLE_FONT, PDF_CID_FONT


def _wrap_pdf_text(text: str, font_name: str, size: float, width: float) -> list[str]:
    lines: list[str] = []
    current = ""
    for char in text:
        candidate = current + char
        if current and pdfmetrics.stringWidth(candidate, font_name, size) > width:
            lines.append(current.rstrip())
            current = char.lstrip()
        else:
            current = candidate
    if current.strip():
        lines.append(current.rstrip())
    return lines


def _draw_pdf_heading(
    pdf: canvas.Canvas,
    section: MarkdownSection,
    *,
    visible_font: str,
    anchor_font: str,
    x: float,
    y: float,
) -> None:
    pdf.setFont(visible_font, 19)
    pdf.setFillColor(HexColor("#163B54"))
    pdf.drawString(x, y, section.title)

    # Invisible text preserves the stable evidence marker for extraction while
    # keeping the reader-facing heading clean.
    pdf.saveState()
    marker = pdf.beginText()
    marker.setTextOrigin(x + pdfmetrics.stringWidth(section.title, visible_font, 19), y)
    marker.setFont(anchor_font, 19)
    marker.setTextRenderMode(3)
    marker.textOut(f" {{#{section.anchor}}}")
    pdf.drawText(marker)
    pdf.restoreState()


def _build_employee_handbook_pdf(source: Path, output: Path) -> None:
    document = _parse_markdown(source)
    visible_font, anchor_font = _register_pdf_fonts()
    page_width, page_height = LETTER
    pdf = canvas.Canvas(
        str(output),
        pagesize=LETTER,
        invariant=1,
        pageCompression=1,
    )
    pdf.setTitle(document.title)
    pdf.setAuthor("Chengming Technology RAG Evaluation")
    pdf.setSubject("Synthetic enterprise knowledge corpus")

    for page_number, section in enumerate(document.sections, start=1):
        margin = 72
        pdf.setFillColor(HexColor("#527384"))
        pdf.setFont(visible_font, 9)
        pdf.drawString(margin, page_height - 48, "澄明科技 · 员工制度参考")
        pdf.drawRightString(page_width - margin, page_height - 48, "内部评测语料")
        pdf.setStrokeColor(HexColor("#C9D7DE"))
        pdf.setLineWidth(0.7)
        pdf.line(margin, page_height - 57, page_width - margin, page_height - 57)

        _draw_pdf_heading(
            pdf,
            section,
            visible_font=visible_font,
            anchor_font=anchor_font,
            x=margin,
            y=page_height - 112,
        )
        pdf.setFillColor(HexColor("#2D5266"))
        pdf.roundRect(margin, page_height - 140, 54, 18, 4, stroke=0, fill=1)
        pdf.setFillColor(HexColor("#FFFFFF"))
        pdf.setFont(visible_font, 8)
        pdf.drawCentredString(margin + 27, page_height - 134.5, "制度要点")

        y = page_height - 176
        pdf.setFont(visible_font, 12)
        pdf.setFillColor(HexColor("#26343B"))
        paragraphs = [item.strip() for item in section.body.split("\n\n") if item.strip()]
        for paragraph in paragraphs:
            for line in _wrap_pdf_text(
                paragraph,
                visible_font,
                12,
                page_width - 2 * margin,
            ):
                pdf.drawString(margin, y, line)
                y -= 22
            y -= 10

        pdf.setStrokeColor(HexColor("#D9E3E8"))
        pdf.line(margin, 54, page_width - margin, 54)
        pdf.setFillColor(HexColor("#6F7E86"))
        pdf.setFont(visible_font, 8)
        pdf.drawString(margin, 38, "企业知识库 RAG 项目 · 虚构教学数据")
        pdf.drawRightString(page_width - margin, 38, f"{page_number} / {len(document.sections)}")
        pdf.showPage()
    pdf.save()


def _set_run_font(
    run,
    *,
    ascii_font: str = DOCX_BODY_FONT,
    east_asia_font: str = DOCX_BODY_FONT,
    size: float | None = None,
    color: str | None = None,
    bold: bool | None = None,
) -> None:
    run.font.name = ascii_font
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), ascii_font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), ascii_font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia_font)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold


def _configure_docx_styles(document: Document) -> None:
    tokens = {
        "Normal": (11, "26343B", 0, 6, 1.25),
        "Heading 1": (16, "2E74B5", 18, 10, 1.0),
        "Heading 2": (13, "2E74B5", 14, 7, 1.0),
        "Heading 3": (12, "1F4D78", 10, 5, 1.0),
    }
    for style_name, (size, color, before, after, line_spacing) in tokens.items():
        style = document.styles[style_name]
        style.font.name = DOCX_BODY_FONT
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), DOCX_BODY_FONT)
        style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), DOCX_BODY_FONT)
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), DOCX_BODY_FONT)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = line_spacing
        style.paragraph_format.keep_with_next = style_name != "Normal"


def _append_page_number(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, text, end])
    _set_run_font(run, size=9, color="6F7E86")


def _add_anchored_heading(document: Document, section: MarkdownSection) -> None:
    paragraph = document.add_paragraph(style="Heading 2")
    visible = paragraph.add_run(section.title)
    _set_run_font(visible, size=13, color="2E74B5", bold=True)
    marker = paragraph.add_run(f" {{#{section.anchor}}}")
    _set_run_font(marker, size=1, color="FFFFFF")
    marker.font.hidden = True


def _build_product_manual_docx(source: Path, output: Path) -> None:
    source_document = _parse_markdown(source)
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.start_type = WD_SECTION.NEW_PAGE

    _configure_docx_styles(document)
    properties = document.core_properties
    properties.title = source_document.title
    properties.subject = "Synthetic enterprise knowledge corpus"
    properties.author = "Chengming Technology RAG Evaluation"
    properties.created = datetime(2000, 1, 1, tzinfo=timezone.utc)
    properties.modified = datetime(2000, 1, 1, tzinfo=timezone.utc)

    header = section.header.paragraphs[0]
    header.text = "澄明科技设备参考"
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header.paragraph_format.space_after = Pt(0)
    for run in header.runs:
        _set_run_font(run, size=9, color="6F7E86", bold=True)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    label = footer.add_run("企业知识库评测语料  ·  ")
    _set_run_font(label, size=9, color="6F7E86")
    _append_page_number(footer)

    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(116)
    kicker = document.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(16)
    _set_run_font(kicker.add_run("PRODUCT REFERENCE MANUAL"), size=10, color="B37A2B", bold=True)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(8)
    _set_run_font(title.add_run("澄明科技设备产品手册"), size=29, color="163B54", bold=True)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(28)
    _set_run_font(subtitle.add_run("ATLAS-X2 星云网关 · ORBIT-M1 极光终端"), size=14, color="527384")

    meta = document.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.paragraph_format.space_after = Pt(96)
    _set_run_font(meta.add_run("规格 · 故障处理 · 工作环境 · 保修政策"), size=10.5, color="6F7E86")

    note = document.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note.paragraph_format.space_after = Pt(0)
    _set_run_font(note.add_run("版本 2026.08  |  虚构教学数据"), size=9.5, color="6F7E86")
    document.add_page_break()

    heading = document.add_paragraph(style="Heading 1")
    _set_run_font(heading.add_run("产品规格与服务政策"), size=16, color="2E74B5", bold=True)
    intro = document.add_paragraph(source_document.introduction)
    intro.paragraph_format.space_after = Pt(12)

    for source_section in source_document.sections:
        _add_anchored_heading(document, source_section)
        paragraph = document.add_paragraph(source_section.body)
        paragraph.paragraph_format.keep_together = True

    document.save(output)


def _extract_pdf_text(path: Path) -> str:
    reader = PyPDF2.PdfReader(str(path))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _extract_docx_text(path: Path) -> str:
    document = Document(path)
    return "\n".join(paragraph.text for paragraph in document.paragraphs)


def _write_manifest(source_outputs: tuple[tuple[Path, Path, str], ...]) -> None:
    documents: list[dict[str, object]] = []
    for source, output, file_format in source_outputs:
        if file_format == "pdf":
            extracted = _extract_pdf_text(output)
        elif file_format == "docx":
            extracted = _extract_docx_text(output)
        else:
            extracted = output.read_text(encoding="utf-8")
        normalized = _normalize_text(extracted)
        documents.append(
            {
                "format": file_format,
                "source_path": source.relative_to(ROOT).as_posix(),
                "output_path": output.relative_to(ROOT).as_posix(),
                "source_sha256": _sha256_bytes(source.read_bytes()),
                "source_normalized_text_sha256": _normalized_sha256(
                    source.read_text(encoding="utf-8"),
                ),
                "extracted_text_sha256": _sha256_bytes(normalized.encode("utf-8")),
                "extracted_text_characters": len(normalized),
            },
        )

    manifest = {
        "schema_version": 1,
        "normalization": "collapse-whitespace-v1",
        "documents": documents,
    }
    (GENERATED_DIR / "manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )


def build() -> None:
    GENERATED_DIR.mkdir(parents=True, exist_ok=True)
    employee_source = SOURCE_DIR / "employee_handbook.md"
    product_source = SOURCE_DIR / "product_manual.md"
    faq_source = SOURCE_DIR / "after_sales_faq.md"
    employee_output = GENERATED_DIR / "employee_handbook.pdf"
    product_output = GENERATED_DIR / "product_manual.docx"
    faq_output = GENERATED_DIR / "after_sales_faq.md"

    _build_employee_handbook_pdf(employee_source, employee_output)
    _build_product_manual_docx(product_source, product_output)
    shutil.copyfile(faq_source, faq_output)
    _write_manifest(
        (
            (employee_source, employee_output, "pdf"),
            (product_source, product_output, "docx"),
            (faq_source, faq_output, "markdown"),
        ),
    )


if __name__ == "__main__":
    build()
